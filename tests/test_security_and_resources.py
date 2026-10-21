from io import BytesIO
import time
import uuid
from docx import Document
from tests.conftest import account,profile,session,send,events,wait_for

def test_auth_rotation_logout_and_bad_tokens(client):
    auth,tokens=account(client)
    assert client.get('/api/v1/users/me').status_code==401
    assert client.get('/api/v1/users/me',headers={'Authorization':'Bearer broken'}).status_code==401
    rotated=client.post('/api/v1/auth/refresh',json={'refreshToken':tokens['refreshToken']})
    assert rotated.status_code==200
    assert client.post('/api/v1/auth/refresh',json={'refreshToken':tokens['refreshToken']}).status_code==401
    new=rotated.json()['data'];headers={'Authorization':'Bearer '+new['accessToken']}
    assert client.get('/api/v1/users/me',headers=headers).status_code==200
    assert client.post('/api/v1/auth/logout',headers=headers).status_code==200
    assert client.get('/api/v1/users/me',headers=auth).status_code==401
    assert client.post('/api/v1/auth/refresh',json={'refreshToken':new['refreshToken']}).status_code==401
    bad=client.post('/api/v1/auth/login',json={'account':'no-user','password':'wrong'}).json()
    assert bad['code']=='AUTH_INVALID_CREDENTIALS' and 'password' not in str(bad.get('details',''))

def test_resume_parse_confirm_snapshot_and_file_access(client,auth):
    document=Document();document.add_paragraph('示例大学 软件工程');document.add_paragraph('Python SQL Git');document.add_paragraph('项目：校园服务应用');document.add_paragraph('英语四级')
    handle=BytesIO();document.save(handle);content=handle.getvalue();key=uuid.uuid4().hex
    def upload():return client.post('/api/v1/resumes',headers={**auth,'Idempotency-Key':key},files={'file':('resume.docx',content,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    response=upload();assert response.status_code==200,response.text
    resume=response.json()['data'];identifier=resume['resumeId']
    assert upload().json()['data']['resumeId']==identifier
    assert client.get(f'/api/v1/resumes/{identifier}/file',headers=auth).content==content
    other,_=account(client)
    assert client.get(f'/api/v1/resumes/{identifier}/file',headers=other).status_code==404
    job=client.post(f'/api/v1/resumes/{identifier}/parse',headers=auth).json()['data']
    assert client.post(f'/api/v1/resumes/{identifier}/parse',headers=auth).json()['data']['jobId']==job['jobId']
    for _ in range(100):
        parsed=client.get(f'/api/v1/resumes/{identifier}/parse',headers=auth).json()['data']
        if parsed['status'] not in ('queued','processing'):break
        time.sleep(.02)
    assert parsed['status']=='review_required'
    assert 'Python' in parsed['skills'] and parsed['education'][0]['school']=='示例大学'
    assert client.patch(f'/api/v1/resumes/{identifier}/parsed-content',headers=auth,json={}).status_code==400
    assert client.post(f'/api/v1/resumes/{identifier}/confirm',headers=auth).status_code==200
    s=session(client,auth,resumeId=identifier)
    assert s['resumeSnapshot']['parsedContent']['skills']==parsed['skills']
    assert client.delete('/api/v1/resumes/'+identifier,headers=auth).status_code==200
    assert client.get('/api/v1/sessions/'+s['sessionId'],headers=auth).json()['data']['resumeSnapshot']['parsedContent']['skills']==parsed['skills']

def test_invalid_upload_and_unconfirmed_resume(client,auth):
    assert client.post('/api/v1/resumes',headers={**auth,'Idempotency-Key':uuid.uuid4().hex},files={'file':('script.exe',b'bad')}).status_code==415
    response=client.post('/api/v1/resumes',headers={**auth,'Idempotency-Key':uuid.uuid4().hex},files={'file':('fake.pdf',b'%PDF-broken')})
    identifier=response.json()['data']['resumeId'];p=profile(client,auth)
    body={'mode':'interview','scene':'employment','sceneConfig':{'targetRole':'后端'},'profileId':p['profileId'],'clientSessionId':uuid.uuid4().hex,'resumeId':identifier}
    assert client.post('/api/v1/sessions',headers={**auth,'Idempotency-Key':uuid.uuid4().hex},json=body).status_code==422
    assert client.post(f'/api/v1/resumes/{identifier}/confirm',headers=auth).status_code==409
    assert client.patch(f'/api/v1/resumes/{identifier}/parsed-content',headers=auth,json={'skills':['Python']}).status_code==200
    assert client.post(f'/api/v1/resumes/{identifier}/confirm',headers=auth).status_code==200
    imported=client.post('/api/v1/profiles/'+p['profileId']+'/resume-import',headers=auth,json={'resumeId':identifier,'selectedFields':['skills']})
    assert imported.status_code==200 and imported.json()['data']['detail']['skills']==['Python']

def test_memory_and_weak_points_are_linked_and_deletable(client,auth):
    s=session(client,auth)['sessionId'];accepted=send(client,auth,s,'我还没有经验，希望多练习。');events(client,auth,s,accepted['responseId'])
    client.post(f'/api/v1/sessions/{s}/finish',headers=auth,json={'reason':'user_end'});wait_for(client,auth,s)
    memories=client.get('/api/v1/users/me/memories',headers=auth).json()['data']['items']
    assert memories and memories[0]['source']['sessionId']==s
    assert client.patch('/api/v1/users/me/memory-settings',headers=auth,json={'enabled':False}).json()['data']['enabled'] is False
    context=client.app.state.worker.context(s)
    assert context['memories']==[] and context['weakPoints']==[]
    body={'dimensionKey':'logic','messageIds':[accepted['userMessageId']],'note':'用STAR补充结果'}
    first=client.post(f'/api/v1/sessions/{s}/weak-points',headers=auth,json=body).json()['data']
    assert first['sessionId']==s
    assert client.post(f'/api/v1/sessions/{s}/weak-points',headers=auth,json=body).json()['data']==first
    assert client.post(f'/api/v1/sessions/{s}/weak-points',headers=auth,json={**body,'messageIds':['somebody-elses-message']}).status_code==422
    assert client.delete('/api/v1/users/me/memories',headers=auth).status_code==200
    assert client.get('/api/v1/users/me/memories',headers=auth).json()['data']['items']==[]
    assert client.get('/api/v1/sessions/'+s,headers=auth).status_code==200
    client.delete('/api/v1/sessions/'+s,headers=auth)
    assert client.get('/api/v1/weak-points',headers=auth).json()['data']['items']==[]
    for table in ('messages','responses','reports','weak_points','memories'):
        assert client.app.state.db.one(f'SELECT COUNT(*) AS n FROM {table} WHERE session_id=?',(s,))['n']==0

def test_openapi_references_and_ui(client):
    schema=client.get('/openapi.json').json()
    assert '/api/v1/sessions' in schema['paths']
    def check(obj):
        if isinstance(obj,dict):
            if '$ref' in obj:
                current=schema
                for key in obj['$ref'][2:].split('/'):current=current[key]
            for value in obj.values():check(value)
        elif isinstance(obj,list):
            for item in obj:check(item)
    check(schema)
    events_schema=schema['paths']['/api/v1/sessions/{session_id}/responses/{response_id}/events']['get']['responses']['200']['content']
    assert 'text/event-stream' in events_schema and 'application/json' not in events_schema
    assert client.get('/').status_code==200
    assert '面试训练台' in client.get('/').text
    assert client.get('/assets/app.js').status_code==200
    response=client.get('/api/v1/does-not-exist')
    assert response.status_code==404 and response.json()['requestId']==response.headers['x-request-id']
