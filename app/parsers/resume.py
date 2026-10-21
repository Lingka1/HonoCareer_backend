"""Conservative local extraction: never invent education or project entries."""
from io import BytesIO
import re
import zipfile
from docx import Document
from pypdf import PdfReader
from app.api.contracts import ParsedContent
from app.services.common import ServiceError

def extract(file_data, file_name):
    if file_name.lower().endswith('.pdf'):
        if not file_data.startswith(b'%PDF-'):
            raise ValueError('文件内容不是有效 PDF')
        reader=PdfReader(BytesIO(file_data))
        if reader.is_encrypted or len(reader.pages)>100:
            raise ValueError('不支持加密或超过 100 页的 PDF')
        text='\n'.join((p.extract_text() or '') for p in reader.pages)
    else:
        with zipfile.ZipFile(BytesIO(file_data)) as archive:
            if sum(x.file_size for x in archive.infolist())>30*1024*1024:
                raise ValueError('Word 解压内容超过 30MB')
            if 'word/document.xml' not in archive.namelist():
                raise ValueError('文件内容不是有效 DOCX')
        doc=Document(BytesIO(file_data))
        text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
    if not text.strip():
        raise ServiceError(422,'PARSE_FAILED','文件未包含可提取文本；扫描件请手动补充或改传可复制文字的 PDF/Word')
    text=text[:100000]
    skills=[x for x in ['Python','Java','JavaScript','TypeScript','SQL','SQLite','MySQL','Redis','Docker','Git','FastAPI','Spring','Vue','React','C++','ArkTS','Linux'] if re.search(r'(?<![\w])'+re.escape(x)+r'(?![\w])',text,re.I)]
    education=[]; projects=[]; certificates=[]
    for line in text.splitlines():
        line=line.strip()
        if not line: continue
        school=re.search(r'([\u4e00-\u9fff]{2,18}(?:大学|学院))',line)
        if school and not any(x['school']==school.group(1) for x in education):
            education.append({'school':school.group(1)})
        if re.match(r'^(?:项目名称|项目)[:：]',line):
            name=re.split('[:：]',line,maxsplit=1)[1].strip()[:120]
            if name: projects.append({'name':name})
        for cert in ['英语四级','英语六级','CET-4','CET-6','普通话二级甲等']:
            if cert in line and cert not in certificates: certificates.append(cert)
    return ParsedContent(education=education[:20],projects=projects[:30],skills=skills,certificates=certificates,warnings=['本地提取仅识别明确文本，请核对学校、项目、时间及技能后确认']).model_dump()
